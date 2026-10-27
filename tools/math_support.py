from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "Introducao_e_Fundamentacao_LabMotion.md"
OUTPUT = ROOT / "outputs" / "LabMotion_Introducao_e_Fundamentacao_Formatado.docx"

BLACK = "161616"
MUTED = "555555"
LIGHT_LINE = "C7CCD1"
BLUE = "B7D3DE"
OFF_WHITE = "F7F3EA"
CORAL = "FF5B39"
LAVENDER = "B7A4EA"
YELLOW = "FBC94A"
LIGHT_BLUE = "A9D3E5"
WHITE = "FFFFFF"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, *, color=BLACK, size=12, inside=True, none=False) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    names = ["top", "left", "bottom", "right"]
    if inside:
        names.extend(["insideH", "insideV"])
    for name in names:
        node = borders.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            borders.append(node)
        node.set(qn("w:val"), "nil" if none else "single")
        if not none:
            node.set(qn("w:sz"), str(size))
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), color)


def set_cell_border(cell, edge: str, *, color=LIGHT_LINE, size=8, value="single") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), value)
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:space"), "0")
    node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_shading(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_keep(paragraph, *, with_next=False, together=False, widow=True) -> None:
    paragraph.paragraph_format.keep_with_next = with_next
    paragraph.paragraph_format.keep_together = together
    paragraph.paragraph_format.widow_control = widow


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def set_run_font(run, name="Arial", size=11, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_language(run, lang="pt-BR") -> None:
    r_pr = run._r.get_or_add_rPr()
    lang_el = r_pr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        r_pr.append(lang_el)
    lang_el.set(qn("w:val"), lang)


def add_badge(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(1.85)
    set_table_borders(table, color=BLACK, size=12, inside=False)
    cell = table.cell(0, 0)
    cell.width = Inches(1.85)
    set_cell_fill(cell, BLUE)
    set_cell_margins(cell, top=70, start=120, bottom=70, end=120)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=8.5, bold=True, color=BLACK)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(24)

    add_badge(doc, "PRODUTO EDUCACIONAL")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("LabMotion")
    set_run_font(run, size=32, bold=True, color=BLACK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("Do dado à lei física")
    set_run_font(run, size=19, bold=False, color=BLACK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    p.paragraph_format.right_indent = Inches(0.55)
    run = p.add_run(
        "Introdução e fundamentação teórica de uma experiência digital investigativa "
        "para o ensino de análise experimental e cinemática."
    )
    set_run_font(run, size=12, color=MUTED)
    p.paragraph_format.line_spacing = 1.25

    trio = doc.add_table(rows=1, cols=3)
    trio.autofit = False
    colors = [YELLOW, LAVENDER, LIGHT_BLUE]
    labels = ["Dados", "Modelos", "Evidências"]
    for idx, (cell, color, label) in enumerate(zip(trio.rows[0].cells, colors, labels)):
        cell.width = Inches(2.05)
        set_cell_fill(cell, color)
        set_cell_margins(cell, top=115, start=90, bottom=115, end=90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p_cell = cell.paragraphs[0]
        p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cell.paragraph_format.space_after = Pt(0)
        run = p_cell.add_run(label)
        set_run_font(run, size=10.5, bold=True)
    set_table_borders(trio, color=BLACK, size=12, inside=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("Física Experimental A • FISC 0099")
    set_run_font(run, size=10.5, bold=True, color=BLACK)

    meta = doc.add_table(rows=4, cols=2)
    meta.autofit = False
    meta.columns[0].width = Inches(1.15)
    meta.columns[1].width = Inches(5.05)
    set_table_borders(meta, inside=False, none=True)
    rows = [
        ("Integrantes", "[preencher nomes]"),
        ("Docente", "[preencher nome]"),
        ("Instituição", "[preencher instituição]"),
        ("Local e ano", "[cidade] • 2026"),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        set_row_cant_split(row)
        for cell in row.cells:
            set_cell_margins(cell, top=110, start=0, bottom=110, end=90)
            set_cell_border(cell, "bottom", color=LIGHT_LINE, size=7)
        p_label = row.cells[0].paragraphs[0]
        p_label.paragraph_format.space_after = Pt(0)
        r_label = p_label.add_run(label)
        set_run_font(r_label, size=9, bold=True)
        p_value = row.cells[1].paragraphs[0]
        p_value.paragraph_format.space_after = Pt(0)
        r_value = p_value.add_run(value)
        set_run_font(r_value, size=9.5, italic=True, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Documento para apresentação e desenvolvimento do produto educacional")
    set_run_font(run, size=8.5, color=MUTED)


def add_footer(section) -> None:
    section.different_first_page_header_footer = True
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.7))
    table.autofit = False
    table.columns[0].width = Inches(5.9)
    table.columns[1].width = Inches(0.8)
    set_table_borders(table, inside=False, none=True)
    p_left = table.cell(0, 0).paragraphs[0]
    p_left.paragraph_format.space_after = Pt(0)
    r = p_left.add_run("LabMotion • Física Experimental A")
    set_run_font(r, size=8.5, color=MUTED)
    p_right = table.cell(0, 1).paragraphs[0]
    p_right.paragraph_format.space_after = Pt(0)
    add_page_number(p_right)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.92)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Inches(0.47)
    normal.paragraph_format.widow_control = True

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Arial"
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    heading1.font.size = Pt(15.5)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor.from_string(BLACK)
    heading1.paragraph_format.space_before = Pt(12)
    heading1.paragraph_format.space_after = Pt(10)
    heading1.paragraph_format.keep_with_next = True
    heading1.paragraph_format.keep_together = True
    heading1.paragraph_format.widow_control = True

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Arial"
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    heading2.font.size = Pt(12.1)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor.from_string(BLACK)
    heading2.paragraph_format.space_before = Pt(11)
    heading2.paragraph_format.space_after = Pt(6)
    heading2.paragraph_format.keep_with_next = True
    heading2.paragraph_format.keep_together = True
    heading2.paragraph_format.widow_control = True

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.core_properties.title = "LabMotion: introdução e fundamentação teórica"
    doc.core_properties.subject = "Produto educacional para Física Experimental A"
    doc.core_properties.author = "Equipe LabMotion"
    doc.core_properties.keywords = "Física Experimental; cinemática; gamificação; produto educacional"


def _nodes(value):
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]


def m_run(text: str, *, plain=False):
    run = OxmlElement("m:r")
    r_pr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p" if plain else "i")
    r_pr.append(sty)
    run.append(r_pr)
    t = OxmlElement("m:t")
    if text.startswith(" ") or text.endswith(" "):
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def m_bar(nodes):
    bar = OxmlElement("m:bar")
    bar_pr = OxmlElement("m:barPr")
    pos = OxmlElement("m:pos")
    pos.set(qn("m:val"), "top")
    bar_pr.append(pos)
    e = OxmlElement("m:e")
    for node in _nodes(nodes):
        e.append(node)
    bar.extend([bar_pr, e])
    return bar


def m_sub(base, subscript):
    obj = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    for node in _nodes(base):
        e.append(node)
    sub = OxmlElement("m:sub")
    for node in _nodes(subscript):
        sub.append(node)
    obj.extend([e, sub])
    return obj


def m_sup(base, superscript):
    obj = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    for node in _nodes(base):
        e.append(node)
    sup = OxmlElement("m:sup")
    for node in _nodes(superscript):
        sup.append(node)
    obj.extend([e, sup])
    return obj


def m_frac(numerator, denominator):
    frac = OxmlElement("m:f")
    num = OxmlElement("m:num")
    for node in _nodes(numerator):
        num.append(node)
    den = OxmlElement("m:den")
    for node in _nodes(denominator):
        den.append(node)
    frac.extend([num, den])
    return frac


def m_rad(radicand):
    rad = OxmlElement("m:rad")
    rad_pr = OxmlElement("m:radPr")
    deg_hide = OxmlElement("m:degHide")
    deg_hide.set(qn("m:val"), "1")
    rad_pr.append(deg_hide)
    e = OxmlElement("m:e")
    for node in _nodes(radicand):
        e.append(node)
    rad.extend([rad_pr, e])
    return rad


def m_sum(subscript, superscript, expression):
    nary = OxmlElement("m:nary")
    nary_pr = OxmlElement("m:naryPr")
    char = OxmlElement("m:chr")
    char.set(qn("m:val"), "∑")
    lim = OxmlElement("m:limLoc")
    lim.set(qn("m:val"), "undOvr")
    nary_pr.extend([char, lim])
    sub = OxmlElement("m:sub")
    for node in _nodes(subscript):
        sub.append(node)
    sup = OxmlElement("m:sup")
    for node in _nodes(superscript):
        sup.append(node)
    e = OxmlElement("m:e")
    for node in _nodes(expression):
        e.append(node)
    nary.extend([nary_pr, sub, sup, e])
    return nary


def equation_nodes(formula: str):
    normalized = re.sub(r"\s+", "", formula.strip().rstrip(".,"))
    x_bar = lambda: m_bar(m_run("x"))
    x_i = lambda: m_sub(m_run("x"), m_run("i"))
    u_a = lambda: m_sub(m_run("u"), m_run("A"))
    u_b = lambda: m_sub(m_run("u"), m_run("B"))
    u_c = lambda: m_sub(m_run("u"), m_run("c"))
    x_0 = lambda: m_sub(m_run("x"), m_run("0", plain=True))
    v_0 = lambda: m_sub(m_run("v"), m_run("0", plain=True))

    if normalized.startswith(r"x=\bar{x}\pmu_c"):
        return [
            m_run("x"), m_run("=", plain=True), x_bar(),
            m_run("±", plain=True), u_c(),
        ]
    if normalized.startswith(r"\bar{x}=\frac{1}{N}\sum"):
        return [
            x_bar(), m_run("=", plain=True),
            m_frac(m_run("1", plain=True), m_run("N")),
            m_sum([m_run("i"), m_run("=", plain=True), m_run("1", plain=True)], m_run("N"), x_i()),
        ]
    if normalized.startswith(r"u_A(\bar{x})=\sqrt"):
        difference_squared = m_sup(
            [m_run("(", plain=True), x_i(), m_run("−", plain=True), x_bar(), m_run(")", plain=True)],
            m_run("2", plain=True),
        )
        numerator = m_sum(
            [m_run("i"), m_run("=", plain=True), m_run("1", plain=True)],
            m_run("N"),
            difference_squared,
        )
        denominator = [
            m_run("N"), m_run("(", plain=True), m_run("N"),
            m_run("−", plain=True), m_run("1", plain=True), m_run(")", plain=True),
        ]
        return [
            u_a(), m_run("(", plain=True), x_bar(), m_run(")", plain=True),
            m_run("=", plain=True), m_rad(m_frac(numerator, denominator)),
        ]
    if normalized.startswith(r"u_c=\sqrt"):
        return [
            u_c(), m_run("=", plain=True),
            m_rad([
                m_sup(u_a(), m_run("2", plain=True)), m_run("+", plain=True),
                m_sup(u_b(), m_run("2", plain=True)),
            ]),
        ]
    if normalized.startswith("x=x_0+vt"):
        return [m_run("x"), m_run("=", plain=True), x_0(), m_run("+", plain=True), m_run("v"), m_run("t")]
    if normalized.startswith(r"x=x_0+v_0t+\frac{1}{2}at^2"):
        return [
            m_run("x"), m_run("=", plain=True), x_0(), m_run("+", plain=True),
            v_0(), m_run("t"), m_run("+", plain=True),
            m_frac(m_run("1", plain=True), m_run("2", plain=True)),
            m_run("a"), m_sup(m_run("t"), m_run("2", plain=True)),
        ]
    raise ValueError(f"Equação não reconhecida: {formula}")


def add_equation(doc: Document, formula: str, number: int) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(6.0)
    table.columns[1].width = Inches(0.55)
    set_table_borders(table, inside=False, none=True)
    row = table.rows[0]
    set_row_cant_split(row)
    left = row.cells[0]
    right = row.cells[1]
    for cell in (left, right):
        set_cell_margins(cell, top=60, start=40, bottom=80, end=40)

    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    omath_para = OxmlElement("m:oMathPara")
    omath_para_pr = OxmlElement("m:oMathParaPr")
    jc = OxmlElement("m:jc")
    jc.set(qn("m:val"), "centerGroup")
    omath_para_pr.append(jc)
    omath = OxmlElement("m:oMath")
    for node in equation_nodes(formula):
        omath.append(node)
    omath_para.extend([omath_para_pr, omath])
    p._p.append(omath_para)

    p_num = right.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_num.paragraph_format.space_before = Pt(5)
    p_num.paragraph_format.space_after = Pt(0)
    r = p_num.add_run(f"({number})")
    set_run_font(r, size=9.5, color=MUTED)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\\\(.+?\\\)|<https?://[^>]+>)")


INLINE_MATH = {
    r"R^2": "R²",
    r"N": "N",
    r"x_i": "xᵢ",
    r"\bar{x}": "x̄",
    r"x=\bar{x}\pm u_c": "x = x̄ ± u_c",
    r"v_0\approx0": "v₀ ≈ 0",
    r"t": "t",
    r"t^2": "t²",
    r"x": "x",
    r"a/2": "a/2",
    r"a=2A": "a = 2A",
    r"y=Ax+B": "y = Ax + B",
    r"A": "A",
    r"B": "B",
    r"\sqrt{3}": "√3",
    r"x\times t": "x × t",
    r"x\times t^2": "x × t²",
}


def add_inline_markdown(paragraph, text: str, *, base_size=10.7, base_color=BLACK) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=base_size, color=base_color)
            set_language(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=base_color)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, italic=True, color=base_color)
        elif token.startswith(r"\("):
            raw = token[2:-2]
            if raw == r"x=\bar{x}\pm u_c":
                omath = OxmlElement("m:oMath")
                for node in equation_nodes(raw):
                    omath.append(node)
                paragraph._p.append(omath)
                run = None
            else:
                rendered = INLINE_MATH.get(raw, raw.replace(r"\pm", "±").replace("^2", "²").replace("_0", "₀"))
                run = paragraph.add_run(rendered)
                set_run_font(run, name="Cambria Math", size=base_size, italic=False, color=base_color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(8.8, base_size - 0.9), color=MUTED)
        if run is not None:
            set_language(run)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=base_size, color=base_color)
        set_language(run)


def parse_markdown(path: Path):
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        raw = lines[i].strip()
        if raw == "## Sugestões de figuras":
            break
        if not raw:
            i += 1
            continue
        if raw.startswith("# "):
            i += 1
            continue
        if raw.startswith("## "):
            blocks.append(("h1", raw[3:].strip()))
            i += 1
            continue
        if raw.startswith("### "):
            blocks.append(("h2", raw[4:].strip()))
            i += 1
            continue
        if raw == r"\[":
            formula_lines = []
            i += 1
            while i < len(lines) and lines[i].strip() != r"\]":
                formula_lines.append(lines[i].strip())
                i += 1
            blocks.append(("math", " ".join(formula_lines)))
            i += 1
            continue
        blocks.append(("p", raw))
        i += 1
    return blocks


def add_process_figure(doc: Document) -> None:
    p_lead = doc.add_paragraph()
    p_lead.paragraph_format.first_line_indent = Inches(0)
    p_lead.paragraph_format.space_before = Pt(10)
    p_lead.paragraph_format.space_after = Pt(5)
    set_keep(p_lead, with_next=True, together=True)
    run = p_lead.add_run("Percurso formativo")
    set_run_font(run, size=10.5, bold=True)

    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    labels = ["1\nMedir", "2\nRepresentar", "3\nModelar", "4\nInterpretar", "5\nArgumentar"]
    colors = [BLUE, OFF_WHITE, LAVENDER, YELLOW, LIGHT_BLUE]
    for cell, label, color in zip(table.rows[0].cells, labels, colors):
        cell.width = Inches(1.31)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, color)
        set_cell_margins(cell, top=100, start=50, bottom=100, end=50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        first, second = label.split("\n")
        r1 = p.add_run(first + "\n")
        set_run_font(r1, size=7.5, bold=True, color=MUTED)
        r2 = p.add_run(second)
        set_run_font(r2, size=8.8, bold=True)
    set_table_borders(table, color=BLACK, size=10, inside=True)
    set_row_cant_split(table.rows[0])

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    set_keep(caption, together=True)
    run = caption.add_run("Figura 1 – Etapas investigativas que organizam as missões do LabMotion.")
    set_run_font(run, size=8.5, italic=True, color=MUTED)


def add_body(doc: Document, blocks) -> None:
    equation_number = 0
    in_references = False
    process_added = False

    for idx, (kind, content) in enumerate(blocks):
        if kind == "h1":
            if content == "Referências":
                p_break = doc.add_paragraph()
                p_break.paragraph_format.space_after = Pt(0)
                p_break.add_run().add_break(WD_BREAK.PAGE)
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.first_line_indent = Inches(0)
            run = p.add_run(content)
            set_run_font(run, size=15.5, bold=True)
            in_references = content == "Referências"
            continue

        if kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.first_line_indent = Inches(0)
            run = p.add_run(content)
            set_run_font(run, size=12.1, bold=True)
            continue

        if kind == "math":
            equation_number += 1
            add_equation(doc, content, equation_number)
            continue

        if kind == "p":
            p = doc.add_paragraph()
            if in_references:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Inches(-0.3)
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.space_after = Pt(6)
                set_keep(p, together=True, widow=True)
                add_inline_markdown(p, content, base_size=9.5)
            else:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Inches(0.47)
                p.paragraph_format.line_spacing = 1.28
                p.paragraph_format.space_after = Pt(4)
                next_is_math = idx + 1 < len(blocks) and blocks[idx + 1][0] == "math"
                set_keep(p, with_next=next_is_math, widow=True)
                add_inline_markdown(p, content, base_size=10.5)

                if content.startswith("O objetivo geral do LabMotion") and not process_added:
                    add_process_figure(doc)
                    process_added = True


def clean_trailing_empty_paragraphs(doc: Document) -> None:
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        p = doc.paragraphs[-1]._element
        p.getparent().remove(p)


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    doc.add_page_break()
    blocks = parse_markdown(SOURCE)
    add_body(doc, blocks)
    clean_trailing_empty_paragraphs(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
