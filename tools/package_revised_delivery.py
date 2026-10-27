from pathlib import Path
import hashlib
import json
import re
import shutil
import zipfile
from html.parser import HTMLParser
from pypdf import PdfReader
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT / 'LabMotion'
REV = PROJECT / 'revisao_21out'
SHARE = REV / 'PARA_COMPARTILHAR'
GROUP = SHARE / 'Grupo'
shutil.copy2(PROJECT / 'entrega/LabMotion_Apresentacao.pptx', GROUP / 'LabMotion_Apresentacao.pptx')
roteiro = (PROJECT / 'entrega/Roteiro_Apresentacao.md').read_text(encoding='utf-8')
roteiro = roteiro.replace(
    'Abrir previamente o produto no dispositivo de apresentação, conferir o acesso e preparar capturas reais da versão utilizada como apoio caso a conexão falhe.',
    'Baixar e extrair o ZIP. Abrir Jogo/LabMotion.html no navegador do computador e abrir os slides. O HTML funciona sem conexão. Preparar capturas reais da versão utilizada como apoio caso haja um problema no equipamento.'
)
roteiro = roteiro.replace(
    'Slide 2. Explicar o problema educacional e o objetivo do produto.',
    'Slide 2. Explicar o problema educacional, sua relevância para a Engenharia e o objetivo do produto.'
)
roteiro = roteiro.replace(
    'Slide 3. Relacionar aprendizagem ativa, feedback e argumentação às decisões do jogo.',
    'Slide 3. Relacionar aprendizagem ativa, investigação guiada, feedback e argumentação às decisões do jogo.'
)
roteiro = roteiro.replace(
    '“A narrativa apresenta um relatório corrompido.',
    '“Um escape room educacional reúne desafios ligados por uma história e um objetivo comum. No LabMotion, a narrativa apresenta um relatório corrompido.'
)
needle = '\n### Matheus\n'
addition = '\n“Na Engenharia, interpretar medidas e gráficos ajuda a analisar processos e qualidade. Escolhemos o tema porque a disciplina exige ligar dados a modelos. A investigação é guiada: os dados são fornecidos, mas a equipe compara interpretações, justifica a escolha e revê a resposta após o feedback. A conclusão precisa apontar a evidência e explicar sua relação com o modelo físico.”\n'
roteiro = roteiro.replace(needle, addition + needle)
(GROUP / 'Roteiro_Apresentacao.md').write_text(roteiro, encoding='utf-8')


def plain(text):
    text = re.sub(r'^#{1,6}\s*', '', text, flags=re.M)
    return text.replace('**', '').replace('`', '')


(GROUP / 'Roteiro_Apresentacao.txt').write_text(plain(roteiro), encoding='utf-8-sig')
read_first = '''LABMOTION — ENTREGA DE 21 DE OUTUBRO DE 2026

O relatório e o jogo estão prontos para avaliação e apresentação.
Ainda não houve aplicação com alunos: os resultados do relatório são de desenvolvimento e verificação técnica.

PARA A PROFESSORA
Anexar Professora/LabMotion_Projeto_Final_21out.pdf e o arquivo Jogo/LabMotion.html, ou enviar este ZIP completo.

PARA OS COLEGAS
Anexar o ZIP completo como documento no WhatsApp ou por e-mail. Todos devem baixar e extrair.
Na pasta Grupo estão o Word editável, os slides e o roteiro de falas.

PARA APRESENTAR O JOGO
1. Extraia o ZIP no computador.
2. Abra Jogo e dê dois cliques em LabMotion.html para abrir no navegador.
3. Abra os slides da pasta Grupo. Projete a janela ou compartilhe a tela.
4. Demonstre uma questão, o feedback e a justificativa. Não é preciso concluir as cinco missões ao vivo.
O jogo em HTML funciona sem internet e sem instalar programas ou entrar em uma conta.

Envie o arquivo de verdade; um caminho C:/Users/... não abre no computador de outra pessoa.
O site antigo é privado e pode estar em outra versão. A entrega atual é este pacote.

O relatório incorpora os cinco comentários da professora e contém as seções completas exigidas.
Confira CONFERENCIA_E_ENVIO.txt para a comparação com o arquivo do colega e mensagens prontas para copiar.
'''
(SHARE / 'LEIA_PRIMEIRO.txt').write_text(read_first, encoding='utf-8-sig')
guide = (SHARE / 'CONFERENCIA_E_ENVIO.md').read_text(encoding='utf-8')
(SHARE / 'CONFERENCIA_E_ENVIO.txt').write_text(plain(guide), encoding='utf-8-sig')

pdf = SHARE / 'Professora/LabMotion_Projeto_Final_21out.pdf'
docx = GROUP / 'LabMotion_Projeto_Final_21out.docx'
game = SHARE / 'Jogo/LabMotion.html'
original_game = PROJECT / 'jogo_portatil/LabMotion.html'
digest = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
assert digest(game) == digest(original_game), 'O jogo foi alterado durante a cópia'
reader = PdfReader(pdf)
assert len(reader.pages) == 34
pdf_text = '\n'.join(page.extract_text() or '' for page in reader.pages)
for required in ['1 Identificação', '2 Introdução', '3 Objetivos', '4 Materiais e métodos', '5 Resultados e discussão', '6 Conclusões', '7 Bibliografia', 'Apêndice H']:
    assert required in pdf_text, required
document = Document(docx)
assert len(document.inline_shapes) == 7
assert len(document.tables) == 17
assert not any('@@' in p.text for p in document.paragraphs)


class Assets(HTMLParser):
    def __init__(self):
        super().__init__(); self.external = []
    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        for key in ['src', 'poster']:
            value = values.get(key, '')
            if value and not value.startswith(('data:', 'blob:', '#')):
                self.external.append([tag, key, value])
        if tag == 'link' and values.get('rel') in ('stylesheet', 'preload', 'modulepreload'):
            value = values.get('href', '')
            if value and not value.startswith('data:'):
                self.external.append([tag, 'href', value])


assets = Assets(); assets.feed(game.read_text(encoding='utf-8'))
assert not assets.external, assets.external
files = sorted(p for p in SHARE.rglob('*') if p.is_file())
manifest = {
    'versao': '21 de outubro de 2026',
    'pdf_paginas': len(reader.pages),
    'figuras': 7,
    'equacoes_numeradas': 7,
    'tabelas': 17,
    'aplicacao_com_estudantes': False,
    'revisao_visual': 'Todas as 34 páginas; acabamento final nas páginas 1 e 31 verificado novamente',
    'jogo_igual_ao_portatil_testado': True,
    'recursos_html_externos': assets.external,
    'arquivos': [{
        'arquivo': p.relative_to(SHARE).as_posix(),
        'bytes': p.stat().st_size,
        'sha256': digest(p)
    } for p in files]
}
(REV / 'manifesto_final.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
output = REV / 'LabMotion_Para_Enviar_21out.zip'
with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
    for p in files:
        archive.write(p, Path('LabMotion_21out') / p.relative_to(SHARE))
with zipfile.ZipFile(output) as archive:
    assert archive.testzip() is None
    assert len(archive.namelist()) == len(files)
print(json.dumps({'zip': str(output), 'bytes': output.stat().st_size, 'arquivos': len(files), 'pdf_paginas': len(reader.pages)}, ensure_ascii=False))
