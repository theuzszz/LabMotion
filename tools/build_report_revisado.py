"""Build the consolidated October 21 version without replacing older deliveries."""
from pathlib import Path
import importlib.util
import re
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT / 'LabMotion'
REV = PROJECT / 'revisao_21out'
OUT = REV / 'PARA_COMPARTILHAR'
spec = importlib.util.spec_from_file_location('base_report', ROOT / 'tools' / 'build_report.py')
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)


def add_math_token(p, token):
    h = base.helper
    math = OxmlElement('m:oMath')
    if token in ('x̄', 't̄'):
        node = h.m_bar(h.m_run(token[0]))
    elif token in ('uA', 'uB', 'uc'):
        node = h.m_sub(h.m_run('u'), h.m_run(token[1], plain=True))
    elif token in ('xᵢ', 'tᵢ', 'x₀', 'v₀'):
        node = h.m_sub(h.m_run(token[0]), h.m_run('i' if token[1] == 'ᵢ' else '0'))
    else:
        acc = OxmlElement('m:acc')
        props = OxmlElement('m:accPr')
        char = OxmlElement('m:chr'); char.set(qn('m:val'), '̂'); props.append(char)
        value = OxmlElement('m:e'); value.append(h.m_run('x'))
        acc.extend([props, value])
        node = h.m_sub(acc, h.m_run('i'))
    math.append(node)
    p._p.append(math)


def inline(p, text, size=12):
    for chunk in re.split(r'(\*\*.*?\*\*|\*[^*]+\*)', text, flags=re.S):
        bold = chunk.startswith('**')
        italic = chunk.startswith('*') and not bold
        chunk = chunk.strip('*') if chunk.startswith('*') else chunk
        for part in re.split(r'(x̂ᵢ|x̄|t̄|xᵢ|tᵢ|x₀|v₀|\buA\b|\buB\b|\buc\b)', chunk):
            if part in ('x̂ᵢ', 'x̄', 't̄', 'xᵢ', 'tᵢ', 'x₀', 'v₀', 'uA', 'uB', 'uc'):
                add_math_token(p, part)
                continue
            run = p.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.bold = bold; run.italic = italic


base.inline = inline
old_table = base.table


def table(doc, lines):
    old_table(doc, lines)
    tbl = doc.tables[-1]
    ratio = (16 / 2.54) / 6.45
    for col in tbl.columns:
        col.width = int(col.width * ratio)
    for row in tbl.rows:
        for cell in row.cells:
            if cell.width:
                cell.width = int(cell.width * ratio)
    spacer = doc.paragraphs[-1]
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)
    for run in spacer.runs: run.font.size = Pt(1)


def figure(doc, filename, caption, num):
    from PIL import Image
    path = PROJECT / filename
    w, h = Image.open(path).size
    width = min(16 / 2.54, 4.25 * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = re.sub(r'^Figura \d+\s*[—–-]\s*', '', caption)
    props = p._p.xpath('.//wp:docPr')
    if props: props[0].set('descr', cap)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = Cm(0)
    cp.paragraph_format.line_spacing = 1
    cp.paragraph_format.space_after = Pt(10)
    cp.paragraph_format.keep_together = True
    inline(cp, f'Figura {num} — {cap}', 10)


def expanded_source():
    text = (REV / 'projeto_revisado.md').read_text(encoding='utf-8')
    refs = (ROOT / 'referencias_verificadas.md').read_text(encoding='utf-8').split('## Bibliografia consolidada')[1].split('## Correções necessárias')[0].strip()
    local = 'UNIVERSIDADE FEDERAL DO VALE DO SÃO FRANCISCO. Física Experimental A FISC 0099. Contrato didático e materiais das aulas 1 e 2. Material da disciplina fornecido pela Profa. Mariele Pinheiro, 2026.'
    refs = refs.replace('VELDKAMP, Alice;', local + '\n\nVELDKAMP, Alice;')
    aps = (ROOT / 'apendices.md').read_text(encoding='utf-8')
    aps = '\n'.join(aps.splitlines()[2:])
    aps = re.sub(r'^## (Apêndice[^\n]+)', r'# \1', aps, flags=re.M)
    aps = re.sub(r'^### ', r'## ', aps, flags=re.M)
    caps = '\n\n'.join([
        'As imagens mostram a interface revisada em uma demonstração técnica, sem respostas de uma turma.',
        '![Abertura com narrativa e objetivo da investigação. Fonte: captura da versão local do LabMotion em 20 out. 2026.](figuras/tela_abertura.png)',
        'O mapa orienta a sequência. As missões seguintes permanecem bloqueadas até que a anterior seja concluída e justificada.',
        '![Mapa das cinco missões e progressão do relatório. Fonte: captura da versão local do LabMotion.](figuras/tela_mapa.png)',
        'A resposta incorreta produz feedback explicativo e permite nova tentativa. As pistas são graduais; não há pontuação por velocidade.',
        '![Feedback da missão de medidas após uma resposta incorreta. Fonte: demonstração técnica do LabMotion.](figuras/tela_feedback.png)',
        'Os desafios de cinemática apresentam tabelas e equações para que as alternativas sejam discutidas a partir dos dados.',
        '![Missão sobre MRU com dados e interpretação do ajuste. Fonte: captura da versão local do LabMotion.](figuras/tela_mru.png)',
    ])
    verification = (ROOT / 'verificacao_relatorio.md').read_text(encoding='utf-8')
    for marker, value in [('@@referencias', refs), ('@@apendices', aps), ('@@capturas', caps), ('@@verificacao', verification)]:
        text = text.replace(marker, value)
    text = text.replace('rubrica do Apêndice C', 'rubrica do Apêndice D')
    (REV / 'projeto_consolidado_21out.md').write_text(text, encoding='utf-8')
    assert '@@' not in text.replace('@@eq:', '')
    return text


def center(doc, text, size=12, before=0, after=0, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    inline(p, f'**{text}**' if bold else text, size)
    return p


def build():
    doc = Document()
    base.helper.configure_document(doc)
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3); sec.left_margin = Cm(3)
    sec.bottom_margin = Cm(2); sec.right_margin = Cm(2)
    sec.footer_distance = Cm(1)
    footer = sec.footer
    for element in list(footer._element): footer._element.remove(element)
    fp = footer.add_paragraph()
    fp.paragraph_format.first_line_indent = Cm(0)
    base.helper.add_page_number(fp)
    for run in fp.runs:
        run.font.name = 'Times New Roman'; run.font.size = Pt(10)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    for name, size in [('Title', 20), ('Subtitle', 14), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 12)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    for name in ('TOC 1', 'TOC 2'):
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        style = doc.styles[name]
        style.font.name = 'Times New Roman'; style.font.size = Pt(11)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(3)

    center(doc, 'UNIVERSIDADE FEDERAL DO VALE DO SÃO FRANCISCO\nUNIVASF', bold=True, after=28)
    center(doc, 'Alinne de Souza Santos Castro\nMatheus Souza da Silva Leite\nJoão Charles Farias Gomes Junior\nArielly Ferreira Marques de Lima', after=36)
    center(doc, 'LabMotion — do dado à lei física', size=18, bold=True, after=12)
    center(doc, 'Produto educacional para Física Experimental A', size=14, after=30)
    center(doc, 'Relatório de desenvolvimento, roteiro de aplicação e gabaritos\nTurma de Física Experimental A\nProfessora Mariele Pinheiro', after=30)
    center(doc, '21 de outubro de 2026')
    doc.add_page_break()
    center(doc, 'RESUMO', bold=True, after=16)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    inline(p, 'Este trabalho apresenta o desenvolvimento do LabMotion, um escape room educacional digital voltado à revisão de análise experimental e cinemática em Física Experimental A, na UNIVASF. A narrativa propõe recuperar um relatório corrompido por meio de cinco missões sobre medidas e incertezas, gráficos científicos, Movimento Retilíneo Uniforme, Movimento Retilíneo Uniformemente Variado e argumentação. O recurso utiliza desafios de múltipla escolha, pistas graduais, feedback explicativo e registro de justificativas para promover discussão em equipe e investigação guiada. A importância para a formação em Engenharia está na interpretação de dados, na escolha de modelos e na comunicação de conclusões apoiadas em evidências. O produto foi implementado e submetido a verificações técnicas de cálculos, progressão, preservação dos registros e exportação da sessão. A entrega inclui uma versão portátil para navegador, roteiro de aplicação, gabaritos e instrumentos de avaliação. Os resultados apresentados se referem ao desenvolvimento e ao teste do produto; os dados numéricos são didáticos e não constituem uma coleta experimental do grupo. Ainda não houve aplicação com estudantes. A avaliação pedagógica será realizada posteriormente com diagnóstico de equipe, rubrica docente e discussão das limitações do procedimento.')
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    inline(p, '**Palavras-chave:** Física experimental; escape room educacional; investigação guiada; cinemática; argumentação.')
    doc.add_page_break()
    center(doc, 'SUMÁRIO', bold=True, after=16)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    instruction.text = ' TOC \\o "1-1" \\h \\z \\u '
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    toc.add_run()._r.extend([begin, instruction, separate, end])
    doc.add_page_break()

    lines = expanded_source().splitlines()
    i = 0; eq = 0; fig = 0; in_references = False; in_appendices = False; section_label = ''
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line.startswith('#'):
            level = min(3, len(line) - len(line.lstrip('#')))
            label = line.lstrip('#').strip()
            if level == 1:
                in_references = label.startswith('7 ')
                section_label = label
            if level == 1 and label.startswith('Apêndice'): in_appendices = True
            if level == 1 and (label.startswith('7 ') or label.startswith('Apêndice')):
                doc.add_page_break()
            p = doc.add_paragraph(label, style=f'Heading {level}')
            p.paragraph_format.first_line_indent = Cm(0)
            i += 1; continue
        if line.startswith('@@eq:'):
            eq += 1; base.equation(doc, line.split(':', 1)[1], eq); i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i]); i += 1
            table(doc, rows); continue
        pic = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if pic:
            fig += 1; figure(doc, pic[2], pic[1], fig); i += 1; continue
        bullet = re.match(r'^([-*]|\d+\.)\s+(.*)', line)
        if bullet:
            is_bullet = bullet[1] in ('-', '*')
            p = base.para(doc, bullet[2] if is_bullet else bullet[1] + ' ' + bullet[2], style='List Bullet' if is_bullet else None)
            p.paragraph_format.first_line_indent = Cm(-.35)
            p.paragraph_format.left_indent = Cm(.6)
        else:
            p = base.para(doc, line)
            if in_references or 'https://' in line:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(5)
                for run in p.runs: run.font.size = Pt(10.5)
        if in_appendices:
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
        if section_label.startswith('Apêndice E'):
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
        i += 1
    doc.core_properties.title = 'LabMotion — projeto completo revisado em 21 de outubro de 2026'
    doc.core_properties.author = 'Alinne de Souza Santos Castro; Matheus Souza da Silva Leite; João Charles Farias Gomes Junior; Arielly Ferreira Marques de Lima'
    doc.core_properties.subject = 'Produto implementado, desenvolvimento e planejamento de aplicação'
    output = OUT / 'Grupo' / 'LabMotion_Projeto_Final_21out.docx'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)
    print(f'Figuras: {fig}; equações: {eq}; tabelas: {len(doc.tables)}; parágrafos: {len(doc.paragraphs)}')


if __name__ == '__main__':
    build()
